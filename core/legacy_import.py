import openpyxl

from core.database import get_connection, init_db

KUTULAR_SHEET_NAME = "Kutular"
KUTULAR_HEADER_ROW = 4
KUTULAR_FIRST_DATA_ROW = 5

ENRETAG_INVOICE_SHEET_NAME = "Enretag Invoice"
ENRETAG_INVOICE_FIRST_DATA_ROW = 2

FEDEX_INVOICE_RAW_SHEET_NAME = "Fedex Invoice"


def import_kutular(xlsx_path):
    workbook = openpyxl.load_workbook(xlsx_path, data_only=True)
    sheet = workbook[KUTULAR_SHEET_NAME]

    connection = get_connection()
    init_db(connection)
    connection.execute("DELETE FROM boxes")

    row_count = 0
    for row in sheet.iter_rows(min_row=KUTULAR_FIRST_DATA_ROW, max_col=20, values_only=True):
        no = row[0]
        if no is None or no == "Total":
            continue

        connection.execute(
            """
            INSERT INTO boxes (
                ship_date, company, customer_code, country, box_no, shipment_type,
                box_type_size, multi_no, tracking_id, customer_shipping_fee,
                customer_packaging_fee, customer_total_invoice, fedex_service_type,
                fedex_duty_invoice_no, fedex_duty_amount, fedex_shipping_invoice_no,
                fedex_shipping_amount, fedex_total_cost, profit_loss
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                str(row[1]) if row[1] is not None else None,
                row[2],
                row[3],
                row[4],
                row[5],
                row[6],
                row[7],
                row[8],
                str(row[9]) if row[9] is not None else None,
                row[10],
                row[11],
                row[12],
                row[13],
                row[14],
                row[15],
                row[16],
                row[17],
                row[18],
                row[19],
            ),
        )
        row_count += 1

    connection.commit()
    connection.close()
    return row_count


def import_fedex_invoice_raw(xlsx_path):
    """Imports the 'Fedex Invoice' sheet as-is: it has no header row and mixed,
    not-fully-understood column semantics, so columns are kept generic (col1..col20)
    rather than guessing labels we aren't confident about."""
    workbook = openpyxl.load_workbook(xlsx_path, data_only=True)
    sheet = workbook[FEDEX_INVOICE_RAW_SHEET_NAME]

    connection = get_connection()
    init_db(connection)
    connection.execute("DELETE FROM legacy_fedex_invoice_raw")

    placeholders = ", ".join(["?"] * 20)
    columns = ", ".join(f"col{i}" for i in range(1, 21))

    row_count = 0
    for row in sheet.iter_rows(min_row=1, max_col=20, values_only=True):
        if all(v is None for v in row):
            continue
        values = tuple(str(v) if v is not None else None for v in row)
        connection.execute(
            f"INSERT INTO legacy_fedex_invoice_raw ({columns}) VALUES ({placeholders})",
            values,
        )
        row_count += 1

    connection.commit()
    connection.close()
    return row_count


def import_enretag_invoices(xlsx_path):
    workbook = openpyxl.load_workbook(xlsx_path, data_only=True)
    sheet = workbook[ENRETAG_INVOICE_SHEET_NAME]

    connection = get_connection()
    init_db(connection)
    connection.execute("DELETE FROM customer_invoices")

    row_count = 0
    for row in sheet.iter_rows(min_row=ENRETAG_INVOICE_FIRST_DATA_ROW, max_col=17, values_only=True):
        invoice_no = row[0]
        if invoice_no is None:
            continue

        connection.execute(
            """
            INSERT INTO customer_invoices (
                invoice_no, customer, invoice_date, due_date, memo, item,
                tracking_id, item_rate, item_amount, currency, ship_via, shipping_date
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                invoice_no,
                row[1],
                str(row[2]) if row[2] is not None else None,
                str(row[3]) if row[3] is not None else None,
                row[6],
                row[7],
                str(row[9]) if row[9] is not None else None,
                row[10],
                row[11],
                row[12],
                row[13],
                str(row[14]) if row[14] is not None else None,
            ),
        )
        row_count += 1

    connection.commit()
    connection.close()
    return row_count
